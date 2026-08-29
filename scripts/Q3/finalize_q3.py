#!/usr/bin/env python3
"""Generate Q3 final tables, workbook, and frozen claims from round36."""

from __future__ import annotations

import csv
import hashlib
import json
from datetime import datetime
from pathlib import Path
from typing import Any

from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parents[2]
ROUND_DIR = ROOT / "results/Q3/experiments/round36_capacity_repair_formal"
REPORT_DIR = ROOT / "results/Q3/reports"
TABLE_DIR = REPORT_DIR / "tables"
TEMPLATE = ROOT / "2023A题/result3.xlsx"
DECISION_ID = "q3_round36_paper_signoff"


def read_json(path: Path) -> dict[str, Any]:
    return json.loads(path.read_text(encoding="utf-8"))


def read_csv(path: Path) -> list[dict[str, str]]:
    with path.open(encoding="utf-8", newline="") as stream:
        return list(csv.DictReader(stream))


def write_csv(path: Path, rows: list[dict[str, Any]]) -> None:
    with path.open("w", encoding="utf-8", newline="") as stream:
        writer = csv.DictWriter(stream, fieldnames=list(rows[0]))
        writer.writeheader()
        writer.writerows(rows)


def sha256(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as stream:
        for block in iter(lambda: stream.read(1024 * 1024), b""):
            digest.update(block)
    return digest.hexdigest()


def main() -> None:
    metric_path = ROUND_DIR / "metrics/q3_formal.json"
    summary_path = ROUND_DIR / "run_summary.json"
    layout_path = ROUND_DIR / "tables/q3_finalist_layout.csv"
    monthly_path = ROUND_DIR / "tables/q3_finalist_monthly.csv"
    evidence = read_json(metric_path)
    summary = read_json(summary_path)
    formal = evidence["fine"]
    geometry = evidence["geometry"]
    convergence = evidence["convergence"]
    comparison = evidence["comparison"]
    design = evidence["design"]
    layout = read_csv(layout_path)
    monthly = read_csv(monthly_path)

    assert summary["comparison"]["main"] == formal
    assert len(layout) == design["mirror_count"] == 2939
    assert len(monthly) == 12
    assert geometry["status"] == "PASS"
    assert convergence["status"] == "PASS"
    assert comparison["pooled_feasible"] and comparison["all_seeds_feasible"]
    area = sum(float(row["width_m"]) * float(row["height_m"]) for row in layout)
    assert abs(area - formal["total_area_m2"]) < 1e-7

    REPORT_DIR.mkdir(parents=True, exist_ok=True)
    TABLE_DIR.mkdir(parents=True, exist_ok=True)
    table1 = [{
        "月份": int(row["month"]),
        "平均光学效率": float(row["eta_total"]),
        "平均余弦效率": float(row["eta_cos"]),
        "平均阴影遮挡效率": float(row["eta_sb"]),
        "平均截断效率": float(row["eta_trunc"]),
        "单位面积镜面平均输出热功率(kW/m^2)": float(row["unit_area_power_kw_m2"]),
    } for row in monthly]
    table2 = [{
        "年平均光学效率": formal["eta_total"],
        "年平均余弦效率": formal["eta_cos"],
        "年平均阴影遮挡效率": formal["eta_sb"],
        "年平均截断效率": formal["eta_trunc"],
        "年平均输出热功率(MW)": formal["field_power_mw"],
        "单位面积镜面年平均输出热功率(kW/m^2)": formal["unit_area_power_kw_m2"],
    }]
    table3 = [{
        "吸收塔x坐标(m)": design["tower_x_m"],
        "吸收塔y坐标(m)": design["tower_y_m"],
        "定日镜总数": design["mirror_count"],
        "定日镜总面积(m^2)": formal["total_area_m2"],
        "镜宽范围(m)": json.dumps(geometry["width_range_m"]),
        "镜高范围(m)": json.dumps(geometry["height_range_m"]),
        "安装高度范围(m)": json.dumps(geometry["installation_height_range_m"]),
    }]
    write_csv(TABLE_DIR / "q3_table1_monthly.csv", table1)
    write_csv(TABLE_DIR / "q3_table2_annual.csv", table2)
    write_csv(TABLE_DIR / "q3_table3_design.csv", table3)

    workbook = load_workbook(TEMPLATE)
    sheet = workbook.active
    for row_index, mirror in enumerate(layout, start=2):
        sheet.cell(row_index, 1, design["tower_x_m"])
        sheet.cell(row_index, 2, design["tower_y_m"])
        sheet.cell(row_index, 3, int(mirror["heliostat_id"]))
        sheet.cell(row_index, 4, float(mirror["width_m"]))
        sheet.cell(row_index, 5, float(mirror["height_m"]))
        sheet.cell(row_index, 6, float(mirror["x_m"]))
        sheet.cell(row_index, 7, float(mirror["y_m"]))
        sheet.cell(row_index, 8, float(mirror["z_m"]))
    workbook.save(REPORT_DIR / "result3.xlsx")

    frozen_at = datetime.now().astimezone().isoformat(timespec="seconds")
    claims: list[dict[str, Any]] = []

    def add(claim_id: str, value: Any, unit: str, source: Path, locator: str) -> None:
        claims.append({
            "claim_id": claim_id, "value": value, "unit": unit,
            "source_file": str(source.relative_to(ROOT)), "source_locator": locator,
            "frozen_at": frozen_at, "frozen_by_skill": "solution-package-builder",
            "decision_id": DECISION_ID,
        })

    for field, unit in {
        "eta_total": "1", "eta_cos": "1", "eta_sb": "1", "eta_trunc": "1",
        "field_power_mw": "MW", "unit_area_power_kw_m2": "kW/m^2",
        "total_area_m2": "m^2",
    }.items():
        add(f"q3_{field}", formal[field], unit, metric_path, f"$.fine.{field}")
    for field, unit in {"tower_x_m": "m", "tower_y_m": "m", "mirror_count": "mirrors"}.items():
        add(f"q3_{field}", design[field], unit, metric_path, f"$.design.{field}")
    for field, unit in {
        "width_range_m": "m", "height_range_m": "m",
        "installation_height_range_m": "m", "minimum_ground_clearance_m": "m",
        "spacing_violation_count": "pairs", "height_exceeds_width_count": "mirrors",
    }.items():
        add(f"q3_{field}", geometry[field], unit, metric_path, f"$.geometry.{field}")
    for month_row in monthly:
        month = int(month_row["month"])
        for field, unit in {
            "eta_total": "1", "eta_cos": "1", "eta_sb": "1", "eta_trunc": "1",
            "unit_area_power_kw_m2": "kW/m^2", "field_power_mw": "MW",
        }.items():
            add(f"q3_month_{month:02d}_{field}", float(month_row[field]), unit,
                monthly_path, f"row[month={month}].{field}")
    for field, unit in {
        "annual_field_power_relative_delta": "1",
        "max_monthly_unit_power_relative_delta": "1",
        "fine_seed_field_power_mw_std": "MW", "fine_seed_field_power_mw_range": "MW",
    }.items():
        add(f"q3_{field}", convergence[field], unit, metric_path, f"$.convergence.{field}")
    add("q3_minimum_seed_power_mw", comparison["minimum_seed_power_mw"], "MW", metric_path, "$.comparison.minimum_seed_power_mw")
    add("q3_unit_area_relative_gain_vs_q2", comparison["unit_area_relative_gain"], "1", metric_path, "$.comparison.unit_area_relative_gain")

    sources = [summary_path, metric_path, layout_path, monthly_path]
    frozen = {
        "schema_version": 1, "question": "Q3", "frozen_at": frozen_at,
        "frozen_by_skill": "solution-package-builder", "decision_id": DECISION_ID,
        "claim_scope": "best_verified_feasible_solution_under_recorded_search_budget_not_global_optimum",
        "source_hashes": {str(path.relative_to(ROOT)): sha256(path) for path in sources},
        "claims": claims,
    }
    (REPORT_DIR / "frozen_numbers.json").write_text(
        json.dumps(frozen, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
    print(f"Generated Q3 final answer with {len(layout)} mirrors and {len(claims)} frozen claims")


if __name__ == "__main__":
    main()
