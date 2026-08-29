#!/usr/bin/env python3
"""Generate Q2 final tables, workbook, and frozen claims from round17 only."""

from __future__ import annotations

import csv
import hashlib
import json
from datetime import datetime
from pathlib import Path
from typing import Any

from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parents[2]
ROUND_DIR = ROOT / "results/Q2/experiments/round17_zoned_formal"
REPORT_DIR = ROOT / "results/Q2/reports"
TABLE_DIR = REPORT_DIR / "tables"
TEMPLATE = ROOT / "2023A题/result2.xlsx"
DECISION_ID = "q2_package_signoff"


def read_json(path: Path) -> dict[str, Any]:
    return json.loads(path.read_text(encoding="utf-8"))


def read_csv(path: Path) -> list[dict[str, str]]:
    with path.open(encoding="utf-8", newline="") as stream:
        return list(csv.DictReader(stream))


def write_csv(path: Path, rows: list[dict[str, Any]]) -> None:
    with path.open("w", encoding="utf-8", newline="") as stream:
        writer = csv.DictWriter(stream, fieldnames=list(rows[0]))
        writer.writeheader()
        writer.writerows(rows)


def sha256(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as stream:
        for block in iter(lambda: stream.read(1024 * 1024), b""):
            digest.update(block)
    return digest.hexdigest()


def main() -> None:
    metric_path = ROUND_DIR / "metrics/q2_zoned_formal.json"
    summary_path = ROUND_DIR / "run_summary.json"
    layout_path = ROUND_DIR / "tables/q2_final_layout.csv"
    monthly_path = ROUND_DIR / "tables/q2_final_monthly.csv"
    evidence = read_json(metric_path)
    summary = read_json(summary_path)
    formal = evidence["formal"]
    geometry = evidence["geometry"]
    convergence = evidence["convergence"]
    layout = read_csv(layout_path)
    monthly = read_csv(monthly_path)

    assert summary["formal"] == formal
    assert len(layout) == formal["mirror_count"] == 2983
    assert len(monthly) == 12
    assert geometry["status"] == "PASS"
    assert convergence["status"] == "PASS"
    assert formal["pooled_mean_feasible"] and formal["all_seeds_feasible"]
    area = formal["mirror_count"] * formal["width_m"] * formal["height_m"]
    assert abs(area - formal["total_area_m2"]) < 1e-8

    REPORT_DIR.mkdir(parents=True, exist_ok=True)
    TABLE_DIR.mkdir(parents=True, exist_ok=True)
    table1 = [
        {
            "月份": int(row["month"]),
            "平均光学效率": float(row["eta_total"]),
            "平均余弦效率": float(row["eta_cos"]),
            "平均阴影遮挡效率": float(row["eta_sb"]),
            "平均截断效率": float(row["eta_trunc"]),
            "单位面积镜面平均输出热功率(kW/m^2)": float(row["unit_area_power_kw_m2"]),
        }
        for row in monthly
    ]
    table2 = [{
        "年平均光学效率": formal["eta_total"],
        "年平均余弦效率": formal["eta_cos"],
        "年平均阴影遮挡效率": formal["eta_sb"],
        "年平均截断效率": formal["eta_trunc"],
        "年平均输出热功率(MW)": formal["field_power_mw"],
        "单位面积镜面年平均输出热功率(kW/m^2)": formal["unit_area_power_kw_m2"],
    }]
    table3 = [{
        "吸收塔x坐标(m)": formal["tower_x_m"],
        "吸收塔y坐标(m)": formal["tower_y_m"],
        "定日镜宽度(m)": formal["width_m"],
        "定日镜高度(m)": formal["height_m"],
        "定日镜安装高度(m)": formal["installation_height_m"],
        "定日镜总数": formal["mirror_count"],
        "定日镜总面积(m^2)": formal["total_area_m2"],
    }]
    write_csv(TABLE_DIR / "q2_table1_monthly.csv", table1)
    write_csv(TABLE_DIR / "q2_table2_annual.csv", table2)
    write_csv(TABLE_DIR / "q2_table3_design.csv", table3)

    workbook = load_workbook(TEMPLATE)
    sheet = workbook.active
    for row_index, mirror in enumerate(layout, start=2):
        sheet.cell(row_index, 1, formal["tower_x_m"])
        sheet.cell(row_index, 2, formal["tower_y_m"])
        sheet.cell(row_index, 3, int(mirror["heliostat_id"]))
        sheet.cell(row_index, 4, float(mirror["width_m"]))
        sheet.cell(row_index, 5, float(mirror["height_m"]))
        sheet.cell(row_index, 6, float(mirror["x_m"]))
        sheet.cell(row_index, 7, float(mirror["y_m"]))
        sheet.cell(row_index, 8, float(mirror["z_m"]))
    workbook.save(REPORT_DIR / "result2.xlsx")

    frozen_at = datetime.now().astimezone().isoformat(timespec="seconds")
    claims: list[dict[str, Any]] = []

    def add(claim_id: str, value: Any, unit: str, source: Path, locator: str) -> None:
        claims.append({
            "claim_id": claim_id,
            "value": value,
            "unit": unit,
            "source_file": str(source.relative_to(ROOT)),
            "source_locator": locator,
            "frozen_at": frozen_at,
            "frozen_by_skill": "solution-package-builder",
            "decision_id": DECISION_ID,
        })

    formal_units = {
        "tower_x_m": "m", "tower_y_m": "m", "width_m": "m", "height_m": "m",
        "installation_height_m": "m", "mirror_count": "mirrors", "total_area_m2": "m^2",
        "minimum_spacing_m": "m", "unpaired_mirror_count": "mirrors", "eta_total": "1",
        "eta_cos": "1", "eta_sb": "1", "eta_trunc": "1", "field_power_mw": "MW",
        "unit_area_power_kw_m2": "kW/m^2", "power_margin_mw": "MW",
        "minimum_fine_seed_power_mw": "MW",
    }
    for field, unit in formal_units.items():
        add(f"q2_{field}", formal[field], unit, metric_path, f"$.formal.{field}")
    for row in monthly:
        month = int(row["month"])
        for field, unit in {
            "eta_total": "1", "eta_cos": "1", "eta_sb": "1", "eta_trunc": "1",
            "unit_area_power_kw_m2": "kW/m^2", "field_power_mw": "MW",
        }.items():
            add(f"q2_month_{month:02d}_{field}", float(row[field]), unit,
                monthly_path, f"row[month={month}].{field}")
    add("q2_convergence_power_relative_delta", convergence["annual_field_power_relative_delta"], "1", metric_path, "$.convergence.annual_field_power_relative_delta")
    add("q2_seed_power_std", convergence["fine_seed_field_power_mw_std"], "MW", metric_path, "$.convergence.fine_seed_field_power_mw_std")
    add("q2_seed_power_range", convergence["fine_seed_field_power_mw_range"], "MW", metric_path, "$.convergence.fine_seed_field_power_mw_range")
    sources = [summary_path, metric_path, layout_path, monthly_path]
    frozen = {
        "schema_version": 1,
        "question": "Q2",
        "frozen_at": frozen_at,
        "frozen_by_skill": "solution-package-builder",
        "decision_id": DECISION_ID,
        "claim_scope_decision_id": "q2_baseline_scope_waiver",
        "source_hashes": {str(path.relative_to(ROOT)): sha256(path) for path in sources},
        "claims": claims,
    }
    (REPORT_DIR / "frozen_numbers.json").write_text(
        json.dumps(frozen, ensure_ascii=False, indent=2) + "\n", encoding="utf-8"
    )
    print(f"Generated Q2 final answer with {len(layout)} mirrors and {len(claims)} frozen claims")


if __name__ == "__main__":
    main()
